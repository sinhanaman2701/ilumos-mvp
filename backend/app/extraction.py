from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any
from difflib import SequenceMatcher


class ExtractionError(RuntimeError):
    """Raised when a chart cannot be converted into a usable draft table."""


@dataclass(frozen=True)
class ExtractionResult:
    rows: list[dict[str, Any]]
    source_text: str
    assets: list[dict[str, Any]]


CLAIM_HEADERS = {
    "patent claim",
    "claim",
    "claim element",
    "parent claim element",
    "challenged claims",
    "challenged claims of the",
}
EVIDENCE_HEADERS = {
    "evidence of infringement",
    "evidence of infringement accused product",
    "accused product feature",
    "accused product feature evidence",
    "evidence",
}


def _normalized(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\x00", "").strip()


def _join_cells(values: list[Any]) -> str:
    return "\n".join(text for value in values if (text := _cell_text(value)))


def _header_role(value: Any) -> str | None:
    normalized = _normalized(value)
    if not normalized:
        return None
    if normalized in CLAIM_HEADERS or "claim element" in normalized or "challenged claim" in normalized:
        return "claim"
    if normalized in EVIDENCE_HEADERS or "evidence of infringement" in normalized:
        return "evidence"
    if "accused product" in normalized and "evidence" in normalized:
        return "evidence"
    return None


def _is_header_noise(claim: str, evidence: str) -> bool:
    claim_key = _normalized(claim)
    evidence_key = _normalized(evidence)
    return (
        claim_key in {"patent", "claims", "claim", "patent claim"} and not evidence
    ) or claim_key.startswith("challenged claims") or evidence_key.startswith("evidence of infringement")


def _find_header(matrix: list[list[Any]]) -> tuple[int, int, int] | None:
    for row_index, row in enumerate(matrix[:12]):
        roles = [_header_role(value) for value in row]
        claim_index = next((index for index, role in enumerate(roles) if role == "claim"), None)
        evidence_index = next((index for index, role in enumerate(roles) if role == "evidence"), None)
        if claim_index is not None and evidence_index is not None:
            return row_index, claim_index, evidence_index
    return None


def _make_row(
    *,
    row_key: str,
    source: dict[str, Any],
    claim: str,
    evidence: str,
    asset_keys: list[str] | None = None,
) -> dict[str, Any] | None:
    if not claim and not evidence:
        return None
    return {
        "draftRowId": row_key,
        "patentClaim": claim,
        "evidenceOfInfringement": evidence,
        "aiEvidence": "",
        "source": source,
        "assetKeys": asset_keys or [],
    }


def _rows_from_matrix(
    matrix: list[list[Any]],
    *,
    filename: str,
    source_file_id: str,
    sheet_name: str,
    row_sources: dict[int, dict[str, Any]],
    asset_keys_by_row: dict[int, list[str]] | None = None,
) -> list[dict[str, Any]]:
    if not matrix:
        return []
    header = _find_header(matrix)
    if header:
        header_row, claim_index, evidence_index = header
        data_start = header_row + 1
    else:
        claim_index = 0
        evidence_index = 1 if len(matrix[0]) > 1 else 0
        data_start = 0

    rows: list[dict[str, Any]] = []
    previous: dict[str, Any] | None = None
    for matrix_index, raw_row in enumerate(matrix[data_start:], start=data_start + 1):
        values = list(raw_row)
        claim = _cell_text(values[claim_index]) if claim_index < len(values) else ""
        if evidence_index < len(values):
            # Keep the evidence column isolated. Columns such as AI notes,
            # comments, or review metadata must not be silently folded into
            # the source evidence field.
            evidence = _cell_text(values[evidence_index])
        else:
            evidence = _join_cells([value for index, value in enumerate(values) if index != claim_index])
        if _header_role(claim) or _header_role(evidence) or _is_header_noise(claim, evidence):
            continue

        source = {
            "fileId": source_file_id,
            "filename": filename,
            "sheet": sheet_name,
            "rowNumber": matrix_index,
            **row_sources.get(matrix_index, {}),
        }
        row = _make_row(
            row_key=f"{source_file_id}:{sheet_name}:{matrix_index}",
            source=source,
            claim=claim,
            evidence=evidence,
            asset_keys=(asset_keys_by_row or {}).get(matrix_index, []),
        )
        if row is None:
            continue
        if not claim and evidence and previous is not None:
            previous["evidenceOfInfringement"] = "\n".join(
                part for part in [previous["evidenceOfInfringement"], evidence] if part
            )
            previous["source"]["continuedRows"] = [
                *previous["source"].get("continuedRows", []),
                matrix_index,
            ]
            previous["assetKeys"] = [
                *previous.get("assetKeys", []),
                *(asset_keys_by_row or {}).get(matrix_index, []),
            ]
            continue
        rows.append(row)
        previous = row
    return rows


def _xlsx_result(data: bytes, filename: str, source_file_id: str) -> ExtractionResult:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    rows: list[dict[str, Any]] = []
    assets: list[dict[str, Any]] = []
    source_lines: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            matrix: list[list[Any]] = []
            row_sources: dict[int, dict[str, Any]] = {}
            asset_keys_by_row: dict[int, list[str]] = {}
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                values = list(row)
                matrix.append(values)
                source_lines.append(f"[SHEET name={worksheet.title} ROW={row_number}]")
                for column_number, value in enumerate(values, start=1):
                    text = _cell_text(value)
                    if text:
                        source_lines.append(
                            f"{worksheet.title}!{get_column_letter(column_number)}{row_number}={text}"
                        )
                last_column = max(len(values), 1)
                row_sources[row_number] = {
                    "cellRange": f"A{row_number}:{get_column_letter(last_column)}{row_number}",
                }

            for image_index, image in enumerate(getattr(worksheet, "_images", []), start=1):
                try:
                    image_data = image._data()
                    anchor = getattr(image, "anchor", None)
                    from_cell = getattr(anchor, "_from", None)
                    anchor_row = int(getattr(from_cell, "row", 0)) + 1
                    anchor_column = int(getattr(from_cell, "col", 0)) + 1
                    extension = str(getattr(image, "format", None) or "png").lower()
                    asset_key = f"{source_file_id}:{worksheet.title}:image:{image_index}"
                    asset_keys_by_row.setdefault(anchor_row, []).append(asset_key)
                    assets.append(
                        {
                            "assetKey": asset_key,
                            "kind": "embedded-image",
                            "filename": f"{filename}-{worksheet.title}-{image_index}.{extension}",
                            "contentType": f"image/{'jpeg' if extension in {'jpg', 'jpeg'} else extension}",
                            "bytes": image_data,
                            "source": {
                                "fileId": source_file_id,
                                "filename": filename,
                                "sheet": worksheet.title,
                                "anchorCell": f"{get_column_letter(anchor_column)}{anchor_row}",
                                "anchorRow": anchor_row,
                            },
                            "candidateRow": f"{source_file_id}:{worksheet.title}:{anchor_row}",
                        }
                    )
                except Exception:
                    continue

            rows.extend(
                _rows_from_matrix(
                    matrix,
                    filename=filename,
                    source_file_id=source_file_id,
                    sheet_name=worksheet.title,
                    row_sources=row_sources,
                    asset_keys_by_row=asset_keys_by_row,
                )
            )
    finally:
        workbook.close()

    if not rows:
        raise ExtractionError(f"No chart rows could be found in {filename}.")
    return ExtractionResult(rows=rows, source_text="\n".join(source_lines), assets=assets)


def _pdf_table_rows(
    page: Any,
    page_number: int,
    table_index: int,
    source_file_id: str,
    filename: str,
) -> tuple[list[dict[str, Any]], list[list[Any]]]:
    finder = getattr(page, "find_tables", None)
    if finder is None:
        return [], []
    try:
        tables = finder().tables
    except Exception:
        return [], []
    if table_index >= len(tables):
        return [], []

    table = tables[table_index]
    try:
        matrix = table.extract()
    except Exception:
        return [], []
    if not matrix:
        return [], []

    header = _find_header(matrix)
    if header:
        header_row, claim_index, evidence_index = header
        # PDF table detectors often place a merged "Patent" header in a
        # different column from the actual claim text. The supplied claim
        # charts use column 1 for the body and column 2 for evidence, while
        # the merged header is reported in columns 2 and 4. If the first
        # header continuation says "Patent", infer the body column from the
        # populated data rows instead of treating every page as one
        # continuation of the first claim.
        header_continuation = (
            _normalized(matrix[header_row + 1][claim_index])
            if header_row + 1 < len(matrix) and claim_index < len(matrix[header_row + 1])
            else ""
        )
        if header_continuation in {"patent", "claim", "claims"}:
            column_scores = {
                column_index: sum(
                    len(_cell_text(row[column_index]))
                    for row in matrix[header_row + 2 :]
                    if column_index < len(row)
                )
                for column_index in range(max(len(row) for row in matrix))
                if column_index != evidence_index
            }
            inferred_claim_index, inferred_score = max(
                column_scores.items(), key=lambda item: item[1], default=(claim_index, 0)
            )
            if inferred_score > 0:
                claim_index = inferred_claim_index
        data_start = header_row + 1
    else:
        claim_index = 0
        evidence_index = max(1, len(matrix[0]) - 1)
        data_start = 0

    rows: list[dict[str, Any]] = []
    previous: dict[str, Any] | None = None
    for row_index, raw_row in enumerate(matrix[data_start:], start=data_start):
        values = list(raw_row)
        claim = _cell_text(values[claim_index]) if claim_index < len(values) else ""
        evidence = _cell_text(values[evidence_index]) if evidence_index < len(values) else ""
        if _header_role(claim) or _header_role(evidence) or _is_header_noise(claim, evidence):
            continue
        row_object = table.rows[row_index] if row_index < len(table.rows) else None
        bbox = list(getattr(row_object, "bbox", table.bbox)) if row_object is not None else list(table.bbox)
        source = {
            "fileId": source_file_id,
            "filename": filename,
            "pageNumber": page_number,
            "tableNumber": table_index + 1,
            "tableRowNumber": row_index + 1,
            "pageBounds": list(page.rect),
            "rowBounds": bbox,
        }
        row = _make_row(
            row_key=f"{source_file_id}:pdf:{page_number}:table:{table_index}:row:{row_index}",
            source=source,
            claim=claim,
            evidence=evidence,
        )
        if row is None:
            continue
        if not claim and evidence and previous is not None:
            previous["evidenceOfInfringement"] = "\n".join(
                part for part in [previous["evidenceOfInfringement"], evidence] if part
            )
            previous["source"]["continuedRows"] = [
                *previous["source"].get("continuedRows", []),
                row_index + 1,
            ]
            continue
        rows.append(row)
        previous = row
    return rows, matrix


def _remove_duplicate_pdf_lead(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Drop a repeated first claim when a PDF contains a cover-page duplicate."""
    if len(rows) < 2:
        return rows
    first_source = rows[0].get("source", {})
    second_source = rows[1].get("source", {})
    if first_source.get("pageNumber") != 1 or second_source.get("pageNumber") != 2:
        return rows

    first_claim = re.findall(r"[a-z0-9]+", rows[0].get("patentClaim", "").lower())
    second_claim = re.findall(r"[a-z0-9]+", rows[1].get("patentClaim", "").lower())
    if len(first_claim) < 8 or len(second_claim) < 8:
        return rows
    if SequenceMatcher(None, first_claim, second_claim).ratio() < 0.72:
        return rows
    return rows[1:]


def _pdf_result(data: bytes, filename: str, source_file_id: str) -> ExtractionResult:
    try:
        import pymupdf
    except ImportError as error:
        raise ExtractionError("PyMuPDF is required for PDF chart import.") from error

    document = pymupdf.open(stream=data, filetype="pdf")
    rows: list[dict[str, Any]] = []
    assets: list[dict[str, Any]] = []
    source_lines: list[str] = []
    try:
        for page_index, page in enumerate(document):
            page_number = page_index + 1
            source_lines.append(f"[PAGE number={page_number}]")
            source_lines.append(page.get_text("text") or "")

            try:
                table_count = len(page.find_tables().tables)
            except Exception:
                table_count = 0
            page_rows: list[dict[str, Any]] = []
            for table_index in range(table_count):
                extracted_rows, matrix = _pdf_table_rows(
                    page,
                    page_number,
                    table_index,
                    source_file_id,
                    filename,
                )
                if extracted_rows:
                    for extracted_row in extracted_rows:
                        if not extracted_row["patentClaim"] and page_rows:
                            page_rows[-1]["evidenceOfInfringement"] = "\n".join(
                                part
                                for part in [
                                    page_rows[-1]["evidenceOfInfringement"],
                                    extracted_row["evidenceOfInfringement"],
                                ]
                                if part
                            )
                            page_rows[-1]["source"]["continuedPages"] = [
                                *page_rows[-1]["source"].get("continuedPages", []),
                                page_number,
                            ]
                        else:
                            page_rows.append(extracted_row)
                    source_lines.extend(
                        f"[TABLE page={page_number} number={table_index + 1}] {row}"
                        for row in matrix
                    )
            for extracted_row in page_rows:
                if not extracted_row["patentClaim"] and rows:
                    rows[-1]["evidenceOfInfringement"] = "\n".join(
                        part
                        for part in [
                            rows[-1]["evidenceOfInfringement"],
                            extracted_row["evidenceOfInfringement"],
                        ]
                        if part
                    )
                    rows[-1]["source"]["continuedPages"] = [
                        *rows[-1]["source"].get("continuedPages", []),
                        page_number,
                    ]
                else:
                    rows.append(extracted_row)

            page_images = page.get_images(full=True)
            row_candidates = [
                (row, row["source"].get("rowBounds"))
                for row in page_rows
                if row["source"].get("rowBounds")
            ]
            for image_index, image_info in enumerate(page_images, start=1):
                try:
                    xref = image_info[0]
                    image = document.extract_image(xref)
                    image_bytes = image["image"]
                    extension = image.get("ext", "png")
                    rects = page.get_image_rects(image_info)
                    rect = rects[0] if rects else None
                    candidate_row = None
                    if rect is not None:
                        page_area = max(float(page.rect.width * page.rect.height), 1.0)
                        image_area = max(float(rect.width * rect.height), 0.0)
                        center_x = (rect.x0 + rect.x1) / 2
                        center_y = (rect.y0 + rect.y1) / 2
                        matches = [
                            row
                            for row, bounds in row_candidates
                            if bounds and bounds[1] <= center_y <= bounds[3] and bounds[0] <= center_x <= bounds[2]
                        ]
                        if image_area / page_area < 0.6 and len(matches) == 1:
                            candidate_row = matches[0]["draftRowId"]
                    asset_key = f"{source_file_id}:pdf:{page_number}:image:{image_index}"
                    assets.append(
                        {
                            "assetKey": asset_key,
                            "kind": "embedded-image",
                            "filename": f"{filename}-page-{page_number}-image-{image_index}.{extension}",
                            "contentType": f"image/{'jpeg' if extension in {'jpg', 'jpeg'} else extension}",
                            "bytes": image_bytes,
                            "source": {
                                "fileId": source_file_id,
                                "filename": filename,
                                "pageNumber": page_number,
                                "imageNumber": image_index,
                                "boundingBox": list(rect) if rect is not None else None,
                            },
                            "candidateRow": candidate_row,
                        }
                    )
                except Exception:
                    continue

            # Vector figures have no embedded image object. Preserve a page render
            # when the page contains drawings but no raster image, without trying
            # to interpret the drawing.
            if not page_images and page.get_drawings():
                try:
                    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
                    asset_key = f"{source_file_id}:pdf:{page_number}:page-render"
                    assets.append(
                        {
                            "assetKey": asset_key,
                            "kind": "page-render",
                            "filename": f"{filename}-page-{page_number}.png",
                            "contentType": "image/png",
                            "bytes": pixmap.tobytes("png"),
                            "source": {
                                "fileId": source_file_id,
                                "filename": filename,
                                "pageNumber": page_number,
                                "pageBounds": list(page.rect),
                            },
                            "candidateRow": None,
                        }
                    )
                except Exception:
                    pass
    finally:
        document.close()

    rows = _remove_duplicate_pdf_lead(rows)
    if not rows:
        raise ExtractionError(
            f"No machine-readable claim-chart rows were found in {filename}. "
            "The original file was preserved, but this chart needs a reviewable import fallback."
        )

    row_by_key = {row["draftRowId"]: row for row in rows}
    for asset in assets:
        candidate_row = asset.get("candidateRow")
        if candidate_row in row_by_key:
            row_by_key[candidate_row].setdefault("assetKeys", []).append(asset["assetKey"])
    return ExtractionResult(rows=rows, source_text="\n".join(source_lines), assets=assets)


def extract_claim_chart(*, filename: str, data: bytes, source_file_id: str) -> ExtractionResult:
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if extension == "xlsx":
        return _xlsx_result(data, filename, source_file_id)
    if extension == "pdf":
        return _pdf_result(data, filename, source_file_id)
    raise ExtractionError(f"Unsupported claim-chart format: {filename}")
